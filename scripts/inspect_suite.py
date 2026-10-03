import openpyxl

wb = openpyxl.load_workbook('NSE_AI_Agent_785_Natural_Language_Test_Cases(1).xlsx', data_only=True)
sheet = wb['All Natural Language Tests']
samples_by_cat = {}
for r in range(2, sheet.max_row + 1):
    cat = sheet.cell(r, 2).value
    if cat not in samples_by_cat:
        samples_by_cat[cat] = []
    if len(samples_by_cat[cat]) < 2:
        samples_by_cat[cat].append({
            'id': sheet.cell(r, 1).value,
            'q': sheet.cell(r, 3).value,
            'intent': sheet.cell(r, 4).value,
            'entity': sheet.cell(r, 5).value,
            'behavior': sheet.cell(r, 6).value,
            'priority': sheet.cell(r, 7).value,
        })

for cat, rows in samples_by_cat.items():
    print(f'=== {cat} ===')
    for row in rows:
        q = row['q']
        it = row['intent']
        ent = row['entity']
        print(f"  {row['id']}: Q='{q}' | ExpIntent='{it}' | ExpEnt='{ent}'")
