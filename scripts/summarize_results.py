import openpyxl

wb = openpyxl.load_workbook('NSE_AI_Agent_Natural_Language_Test_Results.xlsx', data_only=True)
ws_all = wb['2. All 785 Results']

cats = {}
sev_counts = {'Critical': 0, 'High': 0, 'Medium': 0, 'Low': 0}
fail_types = {}
failures_by_sev = {'Critical': [], 'High': [], 'Medium': [], 'Low': []}

for r in range(2, ws_all.max_row + 1):
    tid = ws_all.cell(r, 1).value
    cat = ws_all.cell(r, 2).value
    q = ws_all.cell(r, 3).value
    exp_i = ws_all.cell(r, 4).value
    act_i = ws_all.cell(r, 5).value
    exp_e = ws_all.cell(r, 6).value
    act_e = ws_all.cell(r, 7).value
    exp_b = ws_all.cell(r, 8).value
    act_s = ws_all.cell(r, 9).value
    pri = ws_all.cell(r, 10).value
    status = ws_all.cell(r, 12).value
    f_reason = ws_all.cell(r, 13).value
    f_type = ws_all.cell(r, 14).value
    sev = ws_all.cell(r, 15).value

    if cat not in cats:
        cats[cat] = {'total': 0, 'passed': 0, 'failed': 0}
    cats[cat]['total'] += 1
    if status == 'PASSED':
        cats[cat]['passed'] += 1
    else:
        cats[cat]['failed'] += 1
        fail_types[f_type] = fail_types.get(f_type, 0) + 1
        s = sev if sev in failures_by_sev else (pri if pri in failures_by_sev else 'Medium')
        failures_by_sev[s].append({
            'tid': tid,
            'q': q,
            'exp_i': exp_i,
            'act_i': act_i,
            'exp_e': exp_e,
            'act_e': act_e,
            'exp_b': exp_b,
            'act_b': act_s,
            'f_reason': f_reason,
            'f_type': f_type,
        })
        sev_counts[s] += 1

print('=== CATEGORY BREAKDOWN ===')
for c, st in sorted(cats.items()):
    pr = round((st['passed'] / st['total']) * 100, 1)
    print(f'{c:30s} | Total: {st["total"]:3d} | Passed: {st["passed"]:3d} | Failed: {st["failed"]:3d} | Pass Rate: {pr:5.1f}%')

print('\n=== FAILURE SEVERITY COUNTS ===')
print(sev_counts)

print('\n=== FAILURE TYPES ===')
for ft, cnt in sorted(fail_types.items(), key=lambda x: x[1], reverse=True):
    print(f'{ft:30s}: {cnt}')

print('\n=== CRITICAL FAILURES (first 25) ===')
for item in failures_by_sev['Critical'][:25]:
    print(f"[{item['tid']}] ({item['f_type']}) Q: {item['q']}")
    print(f"   Expected: Intent={item['exp_i']} | Entity={item['exp_e']}")
    print(f"   Actual:   Intent={item['act_i']} | Entity={item['act_e']}")
    print(f"   Reason:   {item['f_reason']}")
    print()

print('\n=== HIGH FAILURES (first 10) ===')
for item in failures_by_sev['High'][:10]:
    print(f"[{item['tid']}] ({item['f_type']}) Q: {item['q']}")
    print(f"   Expected: Intent={item['exp_i']} | Entity={item['exp_e']}")
    print(f"   Actual:   Intent={item['act_i']} | Entity={item['act_e']}")
    print(f"   Reason:   {item['f_reason']}")
    print()
